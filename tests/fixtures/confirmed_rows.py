"""Writes a copy of the BTU workbook with a "Verification Status" column and new rows (test fixture
for the final end-to-end demo). Prints the added rows as JSON.

usage: python3 tests/fixtures/confirmed_rows.py <workbook> <out.xlsx>

The column is added after T and set to CONFIRMED on every existing row (the workbook's own note says
they are already verified awards). New rows:
  REC-0901  BTU-T-0001, a skill they don't hold yet           CONFIRMED  -> V2 for BTU-AI-000125
  REC-0902  BTU-T-0002, a new skill                          PENDING    -> held
  REC-0903  BTU-T-0004, a new skill                          REJECTED   -> held
  REC-0904  malformed Student ID "BTU-T-00X1"                CONFIRMED  -> error
  REC-0905  unknown Student ID BTU-T-0999 using BTU-T-0002's email and Passport ID  -> error"""
import json
import sys

import openpyxl

src, out = sys.argv[1:3]
wb = openpyxl.load_workbook(src)
ws = wb["Passport Skill Records"]
values = openpyxl.load_workbook(src, data_only=True)["Passport Skill Records"]
HEADER = 4
status_col = ws.max_column + 1
ws.cell(HEADER, status_col).value = "Verification Status"
rows = [r for r in range(HEADER + 1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("BTU-T-")]
for r in rows:
    ws.cell(r, status_col).value = "CONFIRMED"


def first_row(student):
    return next(r for r in rows if ws.cell(r, 1).value == student)


def new_skill_for(student):
    held = {ws.cell(r, 11).value for r in rows if ws.cell(r, 1).value == student}
    return next(r for r in rows if ws.cell(r, 11).value not in held)


added = []


def add(record, student, status, student_id=None):
    last = max(r for r in range(HEADER + 1, ws.max_row + 1) if ws.cell(r, 10).value)
    ws.insert_rows(last + 1)
    new = last + 1
    who, skill = first_row(student), new_skill_for(student)
    for c in range(1, 10):
        ws.cell(new, c).value = ws.cell(who, c).value
    if student_id:
        ws.cell(new, 1).value = student_id
    ws.cell(new, 10).value = record
    for c in (11, 12, 13, 15):
        ws.cell(new, c).value = ws.cell(skill, c).value
    ws.cell(new, 14).value = values.cell(skill, 14).value
    ws.cell(new, 16).value = "2026-10-20"
    ws.cell(new, 17).value = "NOT SYNCED"
    ws.cell(new, status_col).value = status
    added.append({"sheet_row": new, "record_id": record, "student_id": ws.cell(new, 1).value,
                  "skill": ws.cell(new, 12).value, "verification_status": status})


add("REC-0901", "BTU-T-0001", "CONFIRMED")
add("REC-0902", "BTU-T-0002", "PENDING")
add("REC-0903", "BTU-T-0004", "REJECTED")
add("REC-0904", "BTU-T-0003", "CONFIRMED", student_id="BTU-T-00X1")
add("REC-0905", "BTU-T-0002", "CONFIRMED", student_id="BTU-T-0999")
wb.save(out)
print(json.dumps(added))
