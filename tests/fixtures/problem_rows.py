"""Writes a copy of the BTU workbook with problem rows added (test fixture for the admin area).

usage: python3 tests/fixtures/problem_rows.py <workbook> <out.xlsx>
Adds: a row whose Student ID is malformed (REC-0981), two rows sharing Record ID REC-0982, and
one valid new skill row for BTU-T-0003 (REC-0983), which the caller can import with a HOLD policy."""
import sys

import openpyxl

src, out = sys.argv[1:3]
wb = openpyxl.load_workbook(src)
ws = wb["Passport Skill Records"]
values = openpyxl.load_workbook(src, data_only=True)["Passport Skill Records"]
rows = [r for r in range(5, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("BTU-T-")]
mine = [r for r in rows if ws.cell(r, 1).value == "BTU-T-0003"]
held = {ws.cell(r, 11).value for r in mine}
skill_row = next(r for r in rows if ws.cell(r, 11).value not in held)


def add(record, student_id=None):
    last = max(r for r in range(5, ws.max_row + 1) if ws.cell(r, 10).value)
    ws.insert_rows(last + 1)
    new = last + 1
    for c in range(1, 10):
        ws.cell(new, c).value = ws.cell(mine[0], c).value
    if student_id:
        ws.cell(new, 1).value = student_id
    ws.cell(new, 10).value = record
    for c in (11, 12, 13, 15):
        ws.cell(new, c).value = ws.cell(skill_row, c).value
    ws.cell(new, 14).value = values.cell(skill_row, 14).value
    ws.cell(new, 16).value = "2026-10-20"
    ws.cell(new, 17).value = "NOT SYNCED"


add("REC-0981", "BTU-X-12")
add("REC-0982")
add("REC-0982")
add("REC-0983")
wb.save(out)
