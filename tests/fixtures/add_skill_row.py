"""Writes a copy of the BTU workbook with one extra skill row for a student (test fixture).

usage: python3 tests/fixtures/add_skill_row.py <workbook> <out.xlsx> <Student ID> <Record ID>
The new row copies the student's identity from their first row and takes a real skill definition
from the workbook that the student doesn't hold yet."""
import sys

import openpyxl

src, out, student, record = sys.argv[1:5]
wb = openpyxl.load_workbook(src)
ws = wb["Passport Skill Records"]
values = openpyxl.load_workbook(src, data_only=True)["Passport Skill Records"]
rows = [r for r in range(5, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("BTU-T-")]
mine = [r for r in rows if ws.cell(r, 1).value == student]
held = {ws.cell(r, 11).value for r in mine}
skill_row = next(r for r in rows if ws.cell(r, 11).value not in held)
last = rows[-1]
ws.insert_rows(last + 1)
new = last + 1
for c in range(1, 10):
    ws.cell(new, c).value = ws.cell(mine[0], c).value
ws.cell(new, 10).value = record
for c in (11, 12, 13, 15):
    ws.cell(new, c).value = ws.cell(skill_row, c).value
ws.cell(new, 14).value = values.cell(skill_row, 14).value
ws.cell(new, 16).value = "2026-10-20"
ws.cell(new, 17).value = "NOT SYNCED"
wb.save(out)
print(ws.cell(skill_row, 12).value)
