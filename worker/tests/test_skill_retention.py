"""Stage 1 regression tests: a new passport version never loses a skill that was already awarded.

The live sheet watcher imports with --apply-edits, so every case runs with apply_edits=True (and the
key ones also without it). Case found in the audit: BTU-AI-000133 (BTU-T-0003, six skills) was issued
as V2 with five skills after one of its rows (REC-0048) left the sheet.

Everything runs on a copy of BTU's workbook in a temporary folder, in local-demo issuance mode: no
blockchain, no Google Sheet, no live data.
"""
from __future__ import annotations

import os
import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl

from worker import config
from worker.db import connect
from worker.sources.xlsx import read_workbook
from worker.sync import Options, run_import
from worker.table import OUTPUT_COLUMNS

WORKBOOK = Path(os.environ.get(
    "BTU_WORKBOOK", config.PORTAL_ROOT / "samples" / "demo" / "workbook.xlsx"))
SHEET = "Passport Skill Records"
HEADER_ROW = 4
COL = {h: i + 1 for i, h in enumerate([
    "Student ID", "First Name", "Last Name", "BTU Email (Test)", "Program", "Education Level", "Passport ID (Existing)",
    "First Issued (Existing)", "Student Status", "Record ID", "Skill ID", "Skill Name", "Skill Category",
    "Badge/Stamp Code", "Skill Description", "Date Earned", *OUTPUT_COLUMNS])}
BADGE = {"AI Skill": "BDG-AI-001", "AI Applied Skill": "BDG-AI-002", "AI Advanced Skill": "BDG-AI-003"}

STUDENT, PASSPORT = "BTU-T-0003", "BTU-AI-000133"   # six skill rows in BTU's workbook
TMP_ROOT = Path(tempfile.mkdtemp(prefix="bdp-retention-"))
BASELINE = TMP_ROOT / "baseline"


def setUpModule():
    config.ISSUANCE_MODE = "local-demo"
    config.set_data_dir(BASELINE)
    conn = connect()
    run_import(conn, read_workbook(WORKBOOK))
    conn.close()


def tearDownModule():
    shutil.rmtree(TMP_ROOT, ignore_errors=True)


def rows_of(ws, student_id: str) -> list[int]:
    return [r for r in range(HEADER_ROW + 1, ws.max_row + 1) if ws.cell(r, COL["Student ID"]).value == student_id]


def unheld_skill(ws, student_id: str) -> dict:
    """A real catalog skill from the workbook that this student doesn't hold yet."""
    held = {ws.cell(r, COL["Skill ID"]).value for r in rows_of(ws, student_id)}
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        sid = ws.cell(r, COL["Skill ID"]).value
        if isinstance(sid, str) and sid.startswith("SKL-") and sid not in held:
            k = {h: ws.cell(r, COL[h]).value for h in ("Skill ID", "Skill Name", "Skill Category", "Skill Description")}
            return k | {"Badge/Stamp Code": BADGE[k["Skill Category"]]}
    raise AssertionError("no unheld skill in the workbook")


def add_skill_row(ws, student_id: str, record_id: str, skill: dict | None = None, **changes) -> int:
    """A new award row for this student, copied from their first row, inserted right below it."""
    src = rows_of(ws, student_id)[0]
    vals = {h: ws.cell(src, c).value for h, c in COL.items() if h not in OUTPUT_COLUMNS}
    vals.update(skill or unheld_skill(ws, student_id))
    vals.update({"Record ID": record_id, "Date Earned": "2026-10-20"}, **changes)
    ws.insert_rows(src + 1)
    for h, v in vals.items():
        ws.cell(src + 1, COL[h]).value = v
    return src + 1


class Base(unittest.TestCase):
    def setUp(self):
        self.dir = TMP_ROOT / f"case-{self.id().split('.')[-1]}"
        shutil.copytree(BASELINE, self.dir)
        config.set_data_dir(self.dir)
        config.ISSUANCE_MODE = "local-demo"
        self.conn = connect()

    def tearDown(self):
        self.conn.close()

    def workbook(self, name: str, edit) -> Path:
        """A copy of BTU's workbook with edit(ws) applied; the original is never written."""
        wb = openpyxl.load_workbook(WORKBOOK)
        edit(wb[SHEET])
        out = self.dir / f"{name}.xlsx"
        wb.save(out)
        return out

    def sync(self, path: Path = WORKBOOK, **opts):
        return run_import(self.conn, read_workbook(path), Options(**opts))

    def versions(self):
        return self.conn.execute(
            """SELECT v.* FROM passport_versions v JOIN passports p ON p.id = v.passport_fk
               WHERE p.passport_id = ? ORDER BY v.version_no""", (PASSPORT,)).fetchall()

    def records(self, version) -> set[str]:
        return {r[0] for r in self.conn.execute(
            "SELECT record_id FROM version_source_records WHERE version_fk = ?", (version["id"],))}

    def current(self):
        return next(v for v in self.versions() if v["status"] == "CURRENT")

    def awarded(self) -> set[str]:
        return {r[0] for r in self.conn.execute(
            "SELECT a.record_id FROM skill_awards a JOIN students s ON s.id = a.student_fk WHERE s.student_id = ?",
            (STUDENT,))}

    def assert_never_shrinks(self):
        vs = self.versions()
        for older, newer in zip(vs, vs[1:]):
            self.assertLessEqual(self.records(older), self.records(newer),
                                 f"V{newer['version_no']} lost skills of V{older['version_no']}")
        self.assertEqual(self.records(self.current()), self.awarded(), "current version holds every awarded skill")


class T01DeletedRow(Base):
    """The audit case: one of BTU-T-0003's six rows disappears from the sheet."""

    def test_deleted_row_keeps_its_skill_and_issues_nothing(self):
        before = self.versions()
        self.assertEqual(len(self.records(before[-1])), 6)
        path = self.workbook("deleted-row", lambda ws: ws.delete_rows(rows_of(ws, STUDENT)[3]))
        for apply_edits in (True, False):
            self.sync(path, apply_edits=apply_edits)
            self.assertEqual(len(self.versions()), len(before), f"apply_edits={apply_edits}: no new version")
            self.assertEqual(len(self.records(self.current())), 6)
            self.assert_never_shrinks()

    def test_deleted_row_plus_new_skill_keeps_all_six(self):
        def edit(ws):
            ws.delete_rows(rows_of(ws, STUDENT)[3])
            add_skill_row(ws, STUDENT, "REC-0901")
        self.sync(self.workbook("deleted-plus-new", edit), apply_edits=True)
        cur = self.current()
        self.assertEqual(cur["version_no"], 2)
        self.assertEqual(len(self.records(cur)), 7, "six earlier skills + the new one")
        self.assert_never_shrinks()


class T02EditedRow(Base):
    def test_skill_edited_in_place_keeps_the_issued_skill(self):
        """Same Record ID, different skill: the issued skill stays; nothing is rewritten or dropped."""
        def edit(ws):
            r = rows_of(ws, STUDENT)[1]
            for h, v in unheld_skill(ws, STUDENT).items():
                ws.cell(r, COL[h]).value = v
        self.sync(self.workbook("edited-in-place", edit), apply_edits=True)
        self.assertEqual(len(self.records(self.current())), 6)
        self.assert_never_shrinks()

    def test_skill_edited_under_a_new_record_id_adds_without_dropping(self):
        """What the watcher does for an edited skill: the row gets a new Record ID. The old award must stay."""
        def edit(ws):
            r = rows_of(ws, STUDENT)[1]
            for h, v in (unheld_skill(ws, STUDENT) | {"Record ID": "REC-0902"}).items():
                ws.cell(r, COL[h]).value = v
        self.sync(self.workbook("edited-new-record", edit), apply_edits=True)
        cur = self.current()
        self.assertEqual(cur["version_no"], 2)
        self.assertEqual(len(self.records(cur)), 7)
        self.assert_never_shrinks()

    def test_profile_edit_reissues_with_every_skill(self):
        """A program change on one row issues a new version (apply_edits), still with all six skills."""
        def edit(ws):
            ws.cell(rows_of(ws, STUDENT)[0], COL["Program"]).value = "Computer Science"
        self.sync(self.workbook("profile-edit", edit), apply_edits=True)
        cur = self.current()
        self.assertEqual(cur["version_no"], 2)
        self.assertEqual(len(self.records(cur)), 6)
        self.assert_never_shrinks()


class T03Reimport(Base):
    def test_reimporting_the_same_data_twice_changes_nothing(self):
        path = self.workbook("with-new", lambda ws: add_skill_row(ws, STUDENT, "REC-0903"))
        for apply_edits in (True, False):
            self.sync(path, apply_edits=apply_edits)
        n_versions = len(self.versions())
        n_awards = self.conn.execute("SELECT COUNT(*) FROM skill_awards").fetchone()[0]
        n_candidates = self.conn.execute("SELECT COUNT(*) FROM version_candidates").fetchone()[0]
        for apply_edits in (True, True, False):
            self.sync(path, apply_edits=apply_edits)
        self.assertEqual(len(self.versions()), n_versions)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM skill_awards").fetchone()[0], n_awards)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM version_candidates").fetchone()[0], n_candidates)
        self.assertEqual(len(self.records(self.current())), 7)
        self.assert_never_shrinks()


class T04NewSkill(Base):
    def test_new_confirmed_skill_is_added_and_all_previous_kept(self):
        before = self.records(self.current())
        self.sync(self.workbook("new-skill", lambda ws: add_skill_row(ws, STUDENT, "REC-0904")), apply_edits=True)
        cur = self.current()
        self.assertEqual(cur["version_no"], 2)
        self.assertEqual(self.records(cur), before | {"REC-0904"})
        self.assert_never_shrinks()

    def test_unconfirmed_skill_is_not_added(self):
        """With a confirmation column, a PENDING row never reaches the passport (live mode included)."""
        def edit(ws):
            ws.cell(HEADER_ROW, 21).value = "Verification Status"
            for r in range(HEADER_ROW + 1, ws.max_row + 1):
                if isinstance(ws.cell(r, 1).value, str) and ws.cell(r, 1).value.startswith("BTU-T-"):
                    ws.cell(r, 21).value = "CONFIRMED"
            ws.cell(add_skill_row(ws, STUDENT, "REC-0905"), 21).value = "PENDING"
        self.sync(self.workbook("pending", edit), apply_edits=True)
        self.assertEqual(len(self.versions()), 1)
        self.assertNotIn("REC-0905", self.awarded())


class T05InvalidRow(Base):
    def test_invalid_row_leaves_the_current_version_unchanged(self):
        before = self.current()
        cases = {
            "bad-date": lambda ws: add_skill_row(ws, STUDENT, "REC-0906", **{"Date Earned": "not a date"}),
            "missing-skill-name": lambda ws: add_skill_row(ws, STUDENT, "REC-0907", **{"Skill Name": None}),
            "bad-student-id": lambda ws: add_skill_row(ws, STUDENT, "REC-0908", **{"Student ID": "BTU-X-3"}),
        }
        for name, edit in cases.items():
            self.sync(self.workbook(name, edit), apply_edits=True)
            cur = self.current()
            self.assertEqual(cur["id"], before["id"], f"{name}: the same version stays current")
            self.assertEqual(self.records(cur), self.records(before), name)
            self.assertEqual(len(self.versions()), 1, name)


class T06IdentityPreserved(Base):
    def test_passport_id_and_first_issue_date_survive_every_version(self):
        p0 = self.conn.execute("SELECT * FROM passports WHERE passport_id = ?", (PASSPORT,)).fetchone()
        steps = [
            ("v2", lambda ws: add_skill_row(ws, STUDENT, "REC-0909")),
            ("v3-deleted", lambda ws: (add_skill_row(ws, STUDENT, "REC-0909"), add_skill_row(ws, STUDENT, "REC-0910"),
                                       ws.delete_rows(rows_of(ws, STUDENT)[-1]))),
        ]
        for name, edit in steps:
            self.sync(self.workbook(name, edit), apply_edits=True)
        p = self.conn.execute("SELECT * FROM passports WHERE passport_id = ?", (PASSPORT,)).fetchone()
        self.assertEqual((p["id"], p["passport_id"], p["first_issued_on"]), (p0["id"], p0["passport_id"], p0["first_issued_on"]))
        self.assertEqual(self.current()["version_no"], 3)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM passports WHERE student_fk = ?", (p["student_fk"],)).fetchone()[0], 1)
        self.assert_never_shrinks()


if __name__ == "__main__":
    unittest.main()
