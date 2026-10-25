"""Stage 2 tests, run against the real BTU workbook:

    python3 -m unittest discover -s worker/tests -t . -v

A baseline database is built once (fresh import of the workbook, which generates every V1 PDF
with BTU's template); each test then works on its own copy of that baseline.

These tests are about the import, so they run in local-demo issuance mode, where a generated
version is issued without a blockchain transaction. Issuance in the default blockchain mode
(candidates waiting for registration) is covered by test_stage3.py."""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import sqlite3
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

import openpyxl

from worker import config
from worker.db import connect
from worker.sources.xlsx import SheetFormatError, read_workbook, write_status_copy
from worker.sync import Options, run_import
from worker.table import OUTPUT_COLUMNS

WORKBOOK = Path(os.environ.get(
    "BTU_WORKBOOK", config.PORTAL_ROOT / "samples" / "demo" / "workbook.xlsx"))
SHEET = "Passport Skill Records"
HEADER_ROW = 4
COL = {h: i + 1 for i, h in enumerate([
    "Student ID", "First Name", "Last Name", "BTU Email (Test)", "Program", "Education Level", "Passport ID (Existing)",
    "First Issued (Existing)", "Student Status", "Record ID", "Skill ID", "Skill Name", "Skill Category",
    "Badge/Stamp Code", "Skill Description", "Date Earned", *OUTPUT_COLUMNS])}

BASELINE = None
TMP_ROOT = Path(tempfile.mkdtemp(prefix="bdp-stage2-"))


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def workbook_rows():
    ws = openpyxl.load_workbook(WORKBOOK, data_only=True)[SHEET]
    return [dict(zip(COL, r)) | {"_row": i} for i, r in enumerate(ws.iter_rows(values_only=True), start=1)
            if i > HEADER_ROW and isinstance(r[0], str) and r[0].startswith("BTU-T-")]


BADGE = {"AI Skill": "BDG-AI-001", "AI Applied Skill": "BDG-AI-002", "AI Advanced Skill": "BDG-AI-003"}


def unheld_skill(student_id: str) -> dict:
    """The full catalog definition of a real skill from the workbook that this student doesn't hold yet."""
    rows = workbook_rows()
    held = {r["Skill ID"] for r in rows if r["Student ID"] == student_id}
    k = next(r for r in rows if r["Skill ID"] not in held)
    return {h: k[h] for h in ("Skill ID", "Skill Name", "Skill Category", "Skill Description")} | {
        "Badge/Stamp Code": BADGE[k["Skill Category"]]}


def edited_workbook(name: str, edit) -> Path:
    """A copy of the real workbook with one change applied by edit(ws). The original is never touched."""
    wb = openpyxl.load_workbook(WORKBOOK)
    edit(wb[SHEET])
    out = TMP_ROOT / f"{name}.xlsx"
    wb.save(out)
    return out


def append_row(ws, values: dict) -> int:
    """Adds a record right under the last data row (above the notes block)."""
    last = max(i for i in range(HEADER_ROW + 1, ws.max_row + 1)
               if isinstance(ws.cell(i, 1).value, str) and ws.cell(i, 1).value.startswith("BTU-T-"))
    ws.insert_rows(last + 1)
    for h, v in values.items():
        ws.cell(last + 1, COL[h]).value = v
    return last + 1


def copy_row(ws, src_row: int, **changes) -> dict:
    vals = {h: ws.cell(src_row, c).value for h, c in COL.items() if h not in OUTPUT_COLUMNS}
    vals["Badge/Stamp Code"] = {"AI Skill": "BDG-AI-001", "AI Applied Skill": "BDG-AI-002",
                                "AI Advanced Skill": "BDG-AI-003"}[vals["Skill Category"]]
    vals.update(changes)
    vals["Passport Sync Status"] = "NOT SYNCED"
    return vals


def setUpModule():
    global BASELINE
    config.ISSUANCE_MODE = "local-demo"
    BASELINE = TMP_ROOT / "baseline"
    config.set_data_dir(BASELINE)
    conn = connect()
    run_import(conn, read_workbook(WORKBOOK))
    conn.close()


def tearDownModule():
    shutil.rmtree(TMP_ROOT, ignore_errors=True)


class Base(unittest.TestCase):
    def setUp(self):
        self.dir = TMP_ROOT / f"case-{self.id().split('.')[-1]}"
        shutil.copytree(BASELINE, self.dir)
        config.set_data_dir(self.dir)
        config.ISSUANCE_MODE = "local-demo"
        self.conn = connect()

    def tearDown(self):
        self.conn.close()

    def q(self, sql, *args):
        return self.conn.execute(sql, args).fetchall()

    def one(self, sql, *args):
        return self.conn.execute(sql, args).fetchone()

    def run_import(self, path=WORKBOOK, **opts):
        table = read_workbook(path)
        return table, run_import(self.conn, table, Options(**opts))

    def versions(self, passport_id):
        return self.q("""SELECT v.* FROM passport_versions v JOIN passports p ON p.id = v.passport_fk
                         WHERE p.passport_id = ? ORDER BY v.version_no""", passport_id)

    def contents(self, version_id):
        return {r[0] for r in self.q("SELECT record_id FROM version_source_records WHERE version_fk = ?", version_id)}

    def codes(self, result, row):
        return [i.code for i in result.validation.issues if i.sheet_row == row]


class T01HeaderDetection(Base):
    def test_header_row_is_found_below_the_title_rows(self):
        t = read_workbook(WORKBOOK)
        self.assertEqual(t.sheet_name, SHEET)
        self.assertEqual(t.header_row, 4)
        self.assertEqual(t.columns["Student ID"], 1)
        self.assertEqual(t.columns["Date Earned"], 16)
        self.assertEqual(t.columns["Sync Error / Notes"], 20)
        self.assertEqual(len(t.rows), 60, "60 record rows")
        self.assertEqual(len(t.ignored_rows), 15, "the notes block under the data is ignored, not imported")
        self.assertIsNone(t.confirmation_column)

    def test_header_moves_if_rows_are_added_above(self):
        path = edited_workbook("two-more-title-rows", lambda ws: ws.insert_rows(1, 2))
        self.assertEqual(read_workbook(path).header_row, 6)

    def test_missing_header_stops_the_import(self):
        path = edited_workbook("renamed-header", lambda ws: setattr(ws.cell(HEADER_ROW, COL["Record ID"]), "value", "Rec"))
        with self.assertRaises(SheetFormatError):
            read_workbook(path)

    def test_wrong_sheet_name_stops_the_import(self):
        def rename(ws):
            ws.title = "Sheet1"
        with self.assertRaises(SheetFormatError):
            read_workbook(edited_workbook("renamed-sheet", rename))


class T02MultipleSkillsOneStudent(Base):
    def test_every_row_becomes_one_award_of_one_passport(self):
        rows = workbook_rows()
        self.assertEqual(self.one("SELECT COUNT(*) FROM skill_awards")[0], 60)
        self.assertEqual(self.one("SELECT COUNT(*) FROM students")[0], 18)
        self.assertEqual(self.one("SELECT COUNT(*) FROM passports")[0], 18)
        for sid in {r["Student ID"] for r in rows}:
            expected = {r["Record ID"] for r in rows if r["Student ID"] == sid}
            got = {r[0] for r in self.q("""SELECT a.record_id FROM skill_awards a JOIN students s ON s.id = a.student_fk
                                           WHERE s.student_id = ?""", sid)}
            self.assertEqual(got, expected, sid)
            cur = self.one("""SELECT v.id FROM passport_versions v JOIN passports p ON p.id = v.passport_fk
                              JOIN students s ON s.id = p.student_fk WHERE s.student_id = ? AND v.status = 'CURRENT'""", sid)
            self.assertEqual(self.contents(cur[0]), expected, f"{sid}: one current version with all their records")

    def test_six_skills_render_on_the_right_number_of_pages(self):
        import pymupdf
        v = self.versions("BTU-AI-000133")  # BTU-T-0003, six skill rows
        self.assertEqual(len(v), 1)
        pdf = pymupdf.open(config.STORAGE_DIR / v[0]["pdf_storage_key"])
        # cover + about + 1 skill grid page (10 per page) + 2 supplement pages (5 rows per page)
        self.assertEqual(pdf.page_count, 5)
        data = json.loads(pdf.embfile_get("passport-data.json"))
        self.assertEqual(len(data["skills"]), 6)


class T03EmailToStudentId(Base):
    def test_each_test_email_maps_to_its_student_id(self):
        for r in workbook_rows():
            s = self.one("SELECT student_id, first_name, last_name FROM students WHERE email = ?", r["BTU Email (Test)"])
            self.assertEqual(s["student_id"], r["Student ID"], r["BTU Email (Test)"])

    def test_mapping_is_by_student_id_not_by_name(self):
        # Two different students share the surname Japaridze; each email reaches only its own record.
        a = self.one("SELECT s.student_id, p.passport_id FROM students s JOIN passports p ON p.student_fk = s.id WHERE email = 'btu-t-0004@example.org'")
        b = self.one("SELECT s.student_id, p.passport_id FROM students s JOIN passports p ON p.student_fk = s.id WHERE email = 'btu-t-0021@example.org'")
        self.assertEqual((a["student_id"], a["passport_id"]), ("BTU-T-0004", "BTU-AI-000127"))
        self.assertEqual((b["student_id"], b["passport_id"]), ("BTU-T-0021", "BTU-AI-000139"))

    def test_import_creates_no_passwords(self):
        self.assertEqual(self.one("SELECT COUNT(*) FROM auth_credentials")[0], 0,
                         "possession of an email is not a login; passwords only come from the seed")


class T04PassportIdentityPreserved(Base):
    def test_passport_ids_and_first_issue_dates_come_from_the_sheet(self):
        for r in workbook_rows():
            p = self.one("""SELECT p.passport_id, p.first_issued_on FROM passports p JOIN students s ON s.id = p.student_fk
                            WHERE s.student_id = ?""", r["Student ID"])
            self.assertEqual(p["passport_id"], r["Passport ID (Existing)"])
            self.assertEqual(p["first_issued_on"], str(r["First Issued (Existing)"])[:10])

    def test_a_changed_first_issue_date_is_refused_not_applied(self):
        rows = [r for r in workbook_rows() if r["Student ID"] == "BTU-T-0001"]
        def edit(ws):
            for r in rows:
                ws.cell(r["_row"], COL["First Issued (Existing)"]).value = "2026-01-01"
        before = self.versions("BTU-AI-000125")
        _, res = self.run_import(edited_workbook("changed-first-issued", edit))
        for r in rows:
            self.assertIn("DATABASE_CONFLICT", self.codes(res, r["_row"]))
            self.assertEqual(res.row_status[r["_row"]]["Passport Sync Status"], "FAILED")
        self.assertEqual(self.one("SELECT first_issued_on FROM passports WHERE passport_id = 'BTU-AI-000125'")[0], "2026-09-15")
        self.assertEqual([tuple(v) for v in self.versions("BTU-AI-000125")], [tuple(v) for v in before])

    def test_missing_passport_id_is_pending_not_invented(self):
        src = workbook_rows()[0]
        def edit(ws):
            append_row(ws, copy_row(ws, src["_row"], **{"Student ID": "BTU-T-0099", "First Name": "Test",
                       "Last Name": "Newcomer", "BTU Email (Test)": "btu-t-0099@example.org",
                       "Passport ID (Existing)": None, "First Issued (Existing)": None, "Record ID": "REC-0999"}))
        path = edited_workbook("new-student-no-passport", edit)
        t, res = self.run_import(path)
        row = next(r.sheet_row for r in t.rows if r.values["Record ID"] == "REC-0999")
        self.assertIn("PASSPORT_PENDING", self.codes(res, row))
        self.assertEqual(res.row_status[row]["Passport Sync Status"], "NOT SYNCED")
        self.assertIn("Passport pending", res.row_status[row]["Sync Error / Notes"])
        self.assertIsNone(self.one("""SELECT 1 FROM passports p JOIN students s ON s.id = p.student_fk
                                      WHERE s.student_id = 'BTU-T-0099'"""))
        # Only an explicit option assigns a new number, and it says so.
        _, res2 = self.run_import(path, assign_new_passport_ids=True)
        self.assertEqual(self.one("""SELECT p.passport_id FROM passports p JOIN students s ON s.id = p.student_fk
                                     WHERE s.student_id = 'BTU-T-0099'""")[0], "BTU-AI-000143")
        self.assertIn("PASSPORT_ID_ASSIGNED", [i.code for i in res2.validation.issues])
        self.assertEqual(len(self.versions("BTU-AI-000143")), 1)


class T05DuplicateRecordIds(Base):
    def test_two_rows_with_one_record_id_are_both_refused(self):
        src = workbook_rows()[0]  # REC-0001
        def edit(ws):
            append_row(ws, copy_row(ws, src["_row"], **unheld_skill(src["Student ID"])))
        before = self.versions("BTU-AI-000125")
        t, res = self.run_import(edited_workbook("duplicate-record-id", edit))
        dup_rows = [r.sheet_row for r in t.rows if r.values["Record ID"] == "REC-0001"]
        self.assertEqual(len(dup_rows), 2)
        for row in dup_rows:
            self.assertIn("DUPLICATE_RECORD_ID", self.codes(res, row))
            self.assertEqual(res.row_status[row]["Passport Sync Status"], "FAILED")
        self.assertEqual(self.one("SELECT COUNT(*) FROM skill_awards WHERE record_id = 'REC-0001'")[0], 1)
        self.assertEqual(len(self.versions("BTU-AI-000125")), len(before), "no new version")

    def test_same_skill_twice_for_one_student_is_refused(self):
        src = workbook_rows()[0]
        def edit(ws):
            append_row(ws, copy_row(ws, src["_row"], **{"Record ID": "REC-0900"}))
        t, res = self.run_import(edited_workbook("duplicate-skill", edit))
        row = next(r.sheet_row for r in t.rows if r.values["Record ID"] == "REC-0900")
        self.assertIn("DUPLICATE_SKILL", self.codes(res, row))
        self.assertIsNone(self.one("SELECT 1 FROM skill_awards WHERE record_id = 'REC-0900'"))


class T06RepeatedImports(Base):
    def test_reimport_changes_nothing(self):
        before = self.q("SELECT id, status, pdf_sha256 FROM passport_versions ORDER BY id")
        for _ in range(2):
            _, res = self.run_import()
            self.assertEqual(res.summary["awards_new"], 0)
            self.assertEqual(res.summary["awards_unchanged"], 60)
            self.assertEqual(res.summary["versions_generated"], 0)
        self.assertEqual([tuple(r) for r in self.q("SELECT id, status, pdf_sha256 FROM passport_versions ORDER BY id")],
                         [tuple(r) for r in before])
        self.assertEqual(self.one("SELECT COUNT(*) FROM skill_awards")[0], 60)

    def test_reimporting_our_own_written_copy_changes_nothing(self):
        _, res = self.run_import()
        out = self.dir / "copy.xlsx"
        write_status_copy(WORKBOOK, out, read_workbook(WORKBOOK), res.row_status)
        _, res2 = self.run_import(out)
        self.assertEqual((res2.summary["awards_new"], res2.summary["versions_generated"],
                          res2.summary["awards_changed_after_accept"]), (0, 0, 0))

    def test_editing_an_accepted_row_does_not_change_the_issued_skill(self):
        src = workbook_rows()[0]
        def edit(ws):
            ws.cell(src["_row"], COL["Skill Description"]).value = "Edited later"
        # the description belongs to the catalog definition, so edit the date instead
        def edit_date(ws):
            ws.cell(src["_row"], COL["Date Earned"]).value = "2026-09-01"
        _, res = self.run_import(edited_workbook("edited-date", edit_date))
        self.assertIn("CHANGED_AFTER_ACCEPT", self.codes(res, src["_row"]))
        self.assertEqual(self.one("SELECT date_earned FROM skill_awards WHERE record_id = ?", src["Record ID"])[0],
                         str(src["Date Earned"])[:10])
        self.assertEqual(res.summary["versions_generated"], 0)


class T07InvalidAndConflictingRecords(Base):
    def _single(self, name, **changes):
        src = workbook_rows()[0]
        def edit(ws):
            append_row(ws, copy_row(ws, src["_row"], **{"Record ID": "REC-0950", **unheld_skill(src["Student ID"]), **changes}))
        t, res = self.run_import(edited_workbook(name, edit))
        row = next((r.sheet_row for r in t.rows if r.values["Record ID"] == "REC-0950"), None)
        return res, row

    def test_row_level_errors_are_reported_and_not_processed(self):
        cases = {
            "bad-student-id": ({"Student ID": "BTU-T-12"}, "INVALID_STUDENT_ID"),
            "missing-date": ({"Date Earned": None}, "MISSING_FIELD"),
            "bad-date": ({"Date Earned": "18/09/2026"}, "INVALID_DATE"),
            "future-date": ({"Date Earned": (date.today() + timedelta(days=30)).isoformat()}, "FUTURE_DATE"),
            "bad-category": ({"Skill Category": "AI Expert Skill"}, "INVALID_CATEGORY"),
            "bad-skill-id": ({"Skill ID": "SK-16"}, "INVALID_SKILL_ID"),
            "badge-mismatch": ({"Badge/Stamp Code": "BDG-AI-003"}, "BADGE_MISMATCH"),
        }
        for name, (changes, code) in cases.items():
            with self.subTest(name):
                res, row = self._single(name, **changes)
                self.assertIn(code, self.codes(res, row))
                self.assertEqual(res.row_status[row]["Passport Sync Status"], "FAILED")
                self.assertTrue(res.row_status[row]["Sync Error / Notes"])
                self.assertIsNone(self.one("SELECT 1 FROM skill_awards WHERE record_id = 'REC-0950'"))
                self.assertEqual(self.one("SELECT COUNT(*) FROM row_issues WHERE record_id = 'REC-0950' AND code = ?", code)[0] >= 1, True)

    def test_identity_conflict_stops_every_row_of_that_student(self):
        rows = [r for r in workbook_rows() if r["Student ID"] == "BTU-T-0002"]
        src = rows[0]
        def edit(ws):
            ws.cell(rows[1]["_row"], COL["Last Name"]).value = "Beridze-Other"
            append_row(ws, copy_row(ws, src["_row"], **{"Record ID": "REC-0951", **unheld_skill("BTU-T-0002")}))
        before = self.versions("BTU-AI-000126")
        t, res = self.run_import(edited_workbook("identity-conflict", edit))
        for r in t.rows:
            if r.values["Student ID"] == "BTU-T-0002":
                self.assertIn("IDENTITY_CONFLICT", self.codes(res, r.sheet_row))
        self.assertIsNone(self.one("SELECT 1 FROM skill_awards WHERE record_id = 'REC-0951'"))
        self.assertEqual(len(self.versions("BTU-AI-000126")), len(before))

    def test_wrong_student_id_never_updates_another_passport(self):
        # A new skill row carries student 0001's Student ID but student 0002's email
        src = workbook_rows()[0]
        def edit(ws):
            append_row(ws, copy_row(ws, src["_row"], **{"Record ID": "REC-0952", **unheld_skill("BTU-T-0001"),
                                                        "BTU Email (Test)": "btu-t-0002@example.org"}))
        before = {pid: len(self.versions(pid)) for pid in ("BTU-AI-000125", "BTU-AI-000126")}
        _, res = self.run_import(edited_workbook("wrong-identity", edit))
        self.assertIsNone(self.one("SELECT 1 FROM skill_awards WHERE record_id = 'REC-0952'"))
        self.assertEqual({pid: len(self.versions(pid)) for pid in before}, before)
        self.assertTrue({"IDENTITY_CONFLICT", "EMAIL_CONFLICT"} & {i.code for i in res.validation.issues})

    def test_other_students_keep_processing(self):
        res, row = self._single("one-bad-row", **{"Date Earned": None})
        self.assertEqual(res.summary["rows_valid"], 60)


class T08MissingPassportFiles(Base):
    def test_missing_current_pdf_is_reported_not_replaced(self):
        v = self.versions("BTU-AI-000130")[-1]
        f = config.STORAGE_DIR / v["pdf_storage_key"]
        f.unlink()
        _, res = self.run_import()
        self.assertIn({"passport_id": "BTU-AI-000130", "version": v["version_no"], "problem": "PDF_MISSING"},
                      res.summary["stored_pdf_problems"])
        self.assertEqual(len(self.versions("BTU-AI-000130")), 1, "no substitute document is generated")
        self.assertFalse(f.exists())

    def test_changed_pdf_is_reported(self):
        v = self.versions("BTU-AI-000131")[-1]
        f = config.STORAGE_DIR / v["pdf_storage_key"]
        f.write_bytes(f.read_bytes() + b"\n%changed")
        _, res = self.run_import()
        self.assertIn("PDF_HASH_MISMATCH", [p["problem"] for p in res.summary["stored_pdf_problems"]])


class T09SkillAdditionAndRegeneration(Base):
    def test_new_skill_creates_the_next_version_and_keeps_everything_else(self):
        rows = [r for r in workbook_rows() if r["Student ID"] == "BTU-T-0001"]
        held = {r["Skill ID"] for r in rows}
        new_skill = next(r for r in workbook_rows() if r["Skill ID"] not in held)
        def edit(ws):
            append_row(ws, copy_row(ws, rows[0]["_row"], **{
                "Record ID": "REC-0961", "Skill ID": new_skill["Skill ID"], "Skill Name": new_skill["Skill Name"],
                "Skill Category": new_skill["Skill Category"], "Skill Description": new_skill["Skill Description"],
                "Date Earned": "2026-10-20",
                "Badge/Stamp Code": {"AI Skill": "BDG-AI-001", "AI Applied Skill": "BDG-AI-002",
                                     "AI Advanced Skill": "BDG-AI-003"}[new_skill["Skill Category"]]}))
        old = self.versions("BTU-AI-000125")
        old_file = config.STORAGE_DIR / old[-1]["pdf_storage_key"]
        old_hash = sha(old_file)
        t, res = self.run_import(edited_workbook("new-skill", edit))
        self.assertEqual(res.summary["versions_generated"], 1)
        new = self.versions("BTU-AI-000125")
        self.assertEqual(len(new), len(old) + 1)
        self.assertEqual([v["status"] for v in new], ["SUPERSEDED"] * len(old) + ["CURRENT"])
        latest = new[-1]
        self.assertEqual(self.contents(latest["id"]), {r["Record ID"] for r in rows} | {"REC-0961"})
        self.assertRegex(latest["verification_code"], r"^BTU-[0-9A-HJKMNP-TV-Z]{4}-[0-9A-HJKMNP-TV-Z]{4}-[0-9A-HJKMNP-TV-Z]{4}$")
        self.assertNotIn(latest["verification_code"], [v["verification_code"] for v in old])
        new_file = config.STORAGE_DIR / latest["pdf_storage_key"]
        self.assertEqual(sha(new_file), latest["pdf_sha256"])
        self.assertTrue(old_file.exists() and sha(old_file) == old_hash, "previous PDF kept, byte for byte")
        p = self.one("SELECT passport_id, first_issued_on FROM passports p JOIN students s ON s.id = p.student_fk WHERE s.student_id = 'BTU-T-0001'")
        self.assertEqual(tuple(p), ("BTU-AI-000125", "2026-09-15"))
        import pymupdf
        data = json.loads(pymupdf.open(new_file).embfile_get("passport-data.json"))
        self.assertEqual(data["passport_id"], "BTU-AI-000125")
        self.assertEqual(data["first_issued"], "2026-09-15")
        self.assertEqual({s["record_id"] for s in data["skills"]}, {r["Record ID"] for r in rows} | {"REC-0961"})
        self.assertIn(latest["verification_code"], pymupdf.open(new_file).metadata["keywords"])
        # a second import of the same file adds nothing
        _, again = self.run_import(edited_workbook("new-skill-again", edit))
        self.assertEqual(again.summary["versions_generated"], 0)


class T10SyncStatusAndErrors(Base):
    def test_status_values_follow_what_really_happened(self):
        rows = workbook_rows()
        src = rows[0]
        def edit(ws):
            append_row(ws, copy_row(ws, src["_row"], **{"Record ID": "REC-0970", **unheld_skill(src["Student ID"]), "Date Earned": "bad"}))
        path = edited_workbook("status", edit)
        t, res = self.run_import(path)
        by_rec = {r.values["Record ID"]: r.sheet_row for r in t.rows}
        for r in rows:
            st = res.row_status[by_rec[r["Record ID"]]]
            if r["Passport Sync Status"] == "SYNCED":
                # simulated values in the supplied workbook are kept, and reported
                self.assertEqual((st["Passport Sync Status"], st["Synced Passport Version"]), ("SYNCED", r["Synced Passport Version"]))
            else:
                self.assertEqual(st["Passport Sync Status"], "PENDING BLOCKCHAIN")
                self.assertIsNone(st["Synced Passport Version"], "no version label until officially issued on chain")
                self.assertIsNone(st["Synced At (UTC)"])
                self.assertIn("not registered on a blockchain", st["Sync Error / Notes"])
        self.assertEqual(len(res.preserved), 16)
        bad = res.row_status[by_rec["REC-0970"]]
        self.assertEqual(bad["Passport Sync Status"], "FAILED")
        self.assertIn("Date Earned", bad["Sync Error / Notes"])
        self.assertFalse(any(s["Passport Sync Status"] == "SYNCED" and s not in [res.row_status[by_rec[r["Record ID"]]] for r in rows if r["Passport Sync Status"] == "SYNCED"]
                             for s in res.row_status.values()), "nothing newly marked SYNCED in stage 2")
        # written to a copy; the input file is not modified
        before = sha(path)
        out = self.dir / "out.xlsx"
        write_status_copy(path, out, t, res.row_status)
        self.assertEqual(sha(path), before)
        ws = openpyxl.load_workbook(out)[SHEET]
        self.assertEqual(ws.cell(by_rec["REC-0970"], COL["Passport Sync Status"]).value, "FAILED")
        self.assertEqual(ws.cell(by_rec["REC-0003"], COL["Passport Sync Status"]).value, "PENDING BLOCKCHAIN")
        self.assertEqual(ws.cell(by_rec["REC-0001"], COL["Passport Sync Status"]).value, "SYNCED")
        self.assertTrue(str(ws.cell(by_rec["REC-0001"], COL["Badge/Stamp Code"]).value).startswith("=IF("), "formulas kept")
        # every issue is traceable to its sheet row and Record ID
        self.assertTrue(self.one("SELECT 1 FROM row_issues WHERE record_id = 'REC-0970' AND sheet_row = ?", by_rec["REC-0970"]))

    def test_confirmation_column_holds_unconfirmed_rows(self):
        src = workbook_rows()[0]
        def edit(ws):
            ws.cell(HEADER_ROW, 21).value = "Verification Status"
            for r in workbook_rows():
                ws.cell(r["_row"], 21).value = "CONFIRMED"
            row = append_row(ws, copy_row(ws, src["_row"], **{"Record ID": "REC-0980", **unheld_skill(src["Student ID"])}))
            ws.cell(row, 21).value = "PENDING"
        t, res = self.run_import(edited_workbook("confirmation", edit))
        self.assertEqual(t.confirmation_column, "Verification Status")
        row = next(r.sheet_row for r in t.rows if r.values["Record ID"] == "REC-0980")
        self.assertIn("NOT_CONFIRMED", self.codes(res, row))
        self.assertEqual(res.row_status[row]["Passport Sync Status"], "NOT SYNCED")
        self.assertIsNone(self.one("SELECT 1 FROM skill_awards WHERE record_id = 'REC-0980'"))
        self.assertEqual(res.summary["versions_generated"], 0)


class T12FailedUpdateKeepsPreviousVersion(Base):
    def test_render_failure_leaves_the_current_version_in_place(self):
        src = workbook_rows()[0]
        def edit(ws):
            append_row(ws, copy_row(ws, src["_row"], **{"Record ID": "REC-0990", **unheld_skill(src["Student ID"])}))
        path = edited_workbook("fail-then-retry", edit)
        before = [tuple(v) for v in self.versions("BTU-AI-000125")]
        files_before = sorted(p.name for p in (config.STORAGE_DIR / "passports" / "BTU-AI-000125").iterdir())

        def broken_renderer(*_args, **_kw):
            raise RuntimeError("simulated renderer crash")
        t, res = self.run_import(path, renderer=broken_renderer)
        row = next(r.sheet_row for r in t.rows if r.values["Record ID"] == "REC-0990")
        self.assertEqual([tuple(v) for v in self.versions("BTU-AI-000125")], before, "no version added or changed")
        self.assertEqual(sorted(p.name for p in (config.STORAGE_DIR / "passports" / "BTU-AI-000125").iterdir()), files_before)
        self.assertEqual(res.row_status[row]["Passport Sync Status"], "FAILED")
        self.assertIn("simulated renderer crash", res.row_status[row]["Sync Error / Notes"])
        self.assertIn("remains V1", res.row_status[row]["Sync Error / Notes"])
        self.assertEqual(self.one("SELECT COUNT(*) FROM generation_failures")[0], 1)

        _, retry = self.run_import(path)  # working renderer: the same record is picked up again
        self.assertEqual(retry.summary["versions_generated"], 1)
        self.assertEqual(self.versions("BTU-AI-000125")[-1]["status"], "CURRENT")
        self.assertIn("REC-0990", self.contents(self.versions("BTU-AI-000125")[-1]["id"]))

    def test_database_failure_after_storing_leaves_no_orphan_file(self):
        src = workbook_rows()[0]
        def edit(ws):
            append_row(ws, copy_row(ws, src["_row"], **{"Record ID": "REC-0991", **unheld_skill(src["Student ID"])}))
        from worker import issue
        real = issue.transaction
        class failing_tx(real):
            def __enter__(self):  # the database refuses the write after the PDF was stored
                super().__enter__()
                self.conn.execute("ROLLBACK")
                raise sqlite3.OperationalError("simulated database failure")
        issue.transaction = failing_tx
        try:
            _, res = self.run_import(edited_workbook("db-fail", edit))
        finally:
            issue.transaction = real
        self.assertEqual(res.summary["generation_failures"], 1)
        self.assertEqual(len(self.versions("BTU-AI-000125")), 1)
        self.assertEqual(len(list((config.STORAGE_DIR / "passports" / "BTU-AI-000125").iterdir())), 1)

    def test_history_cannot_be_deleted(self):
        with self.assertRaises(sqlite3.IntegrityError):
            self.conn.execute("DELETE FROM passport_versions")
        with self.assertRaises(sqlite3.IntegrityError):
            self.conn.execute("DELETE FROM skill_awards")


if __name__ == "__main__":
    unittest.main()
