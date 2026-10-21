"""Reads the BTU skill workbook (.xlsx) into a SourceTable, and writes Q–T results to a copy."""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

import openpyxl

from .. import config
from ..table import (INPUT_COLUMNS, OUTPUT_COLUMNS, RECORD_ID, STUDENT_ID, SourceRow, SourceTable)

HEADER_SEARCH_ROWS = 50
ID_LIKE = re.compile(r"^[A-Z]{2,}-[A-Z0-9-]+$")


class SheetFormatError(Exception):
    """The workbook doesn't have the expected sheet or header row. Nothing is imported."""


def _norm_header(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def _empty(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def find_header_row(ws, required=INPUT_COLUMNS) -> tuple[int, dict[str, int]]:
    """The header row is the first row that contains every required column name.
    Title and explanation rows above it are skipped."""
    for r in range(1, min(ws.max_row, HEADER_SEARCH_ROWS) + 1):
        names = {}
        for c in range(1, ws.max_column + 1):
            h = _norm_header(ws.cell(r, c).value)
            if h:
                if h in names:
                    raise SheetFormatError(f"Row {r}: column header {h!r} appears twice")
                names[h] = c
        if all(h in names for h in required):
            return r, names
    raise SheetFormatError(
        f"No header row with all of these columns in the first {HEADER_SEARCH_ROWS} rows: {', '.join(required)}"
    )


def read_workbook(path: str | Path, sheet_name: str | None = None) -> SourceTable:
    path = Path(path)
    sheet_name = sheet_name or config.SHEET_NAME
    blob = path.read_bytes()
    values_wb = openpyxl.load_workbook(path, data_only=True)    # cached results of formulas
    formulas_wb = openpyxl.load_workbook(path, data_only=False)  # the formulas themselves
    if sheet_name not in values_wb.sheetnames:
        raise SheetFormatError(f"Worksheet {sheet_name!r} not found; the workbook has {values_wb.sheetnames}")
    ws, wf = values_wb[sheet_name], formulas_wb[sheet_name]

    header_row, columns = find_header_row(ws)
    missing_out = [h for h in OUTPUT_COLUMNS if h not in columns]
    if missing_out:
        raise SheetFormatError(f"Header row {header_row} lacks the system columns {missing_out}")
    confirmation = config.CONFIRMATION_COLUMN if config.CONFIRMATION_COLUMN in columns else None
    input_headers = INPUT_COLUMNS + ([confirmation] if confirmation else [])

    rows, ignored, blank = [], [], 0
    for r in range(header_row + 1, ws.max_row + 1):
        vals = {h: ws.cell(r, columns[h]).value for h in input_headers}
        outs = {h: ws.cell(r, columns[h]).value for h in OUTPUT_COLUMNS}
        if all(_empty(v) for v in [*vals.values(), *outs.values()]):
            blank += 1
            continue
        first = vals[STUDENT_ID]
        others_empty = all(_empty(v) for h, v in vals.items() if h != STUDENT_ID) and all(_empty(v) for v in outs.values())
        if others_empty and not (isinstance(first, str) and ID_LIKE.match(first.strip())):
            ignored.append((r, f"explanatory text, not a record: {str(first)[:60]!r}"))
            continue
        row = SourceRow(sheet_row=r, values=vals, outputs=outs)
        for h in input_headers:
            f = wf.cell(r, columns[h]).value
            if isinstance(f, str) and f.startswith("=") and _empty(vals[h]):
                row.notes.append(f"{h} is a formula with no calculated value in the file")
        rows.append(row)

    return SourceTable(
        source_type="XLSX", source_name=path.name, source_sha256=hashlib.sha256(blob).hexdigest(),
        sheet_name=sheet_name, header_row=header_row, columns=columns, rows=rows,
        ignored_rows=ignored, blank_rows=blank, confirmation_column=confirmation,
    )


def write_status_copy(source_path: str | Path, out_path: str | Path, table: SourceTable,
                      updates: dict[int, dict[str, object]], extra_cells: dict[tuple[int, str], object] | None = None) -> int:
    """Writes Q–T values (and any extra cells, e.g. a newly assigned Passport ID) into a COPY of the
    workbook. The original file is never modified. Formulas and all other cells are kept.
    Returns the number of cells changed."""
    wb = openpyxl.load_workbook(source_path, data_only=False)
    ws = wb[table.sheet_name]
    changed = 0
    for sheet_row, vals in updates.items():
        for header, value in vals.items():
            cell = ws.cell(sheet_row, table.columns[header])
            new = None if value in ("", None) else value
            if cell.value != new:
                cell.value = new
                changed += 1
    for (sheet_row, header), value in (extra_cells or {}).items():
        cell = ws.cell(sheet_row, table.columns[header])
        if cell.value != value:
            cell.value = value
            changed += 1
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return changed

